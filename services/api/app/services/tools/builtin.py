"""Built-in tools. Code, shell, Git, image and document work runs in the sandbox; file and
web access runs in the worker, through workspace.py and netguard.py."""

import asyncio
import hashlib
import html
import json
import os
import shlex
import shutil
import tempfile
import uuid
from html.parser import HTMLParser
from pathlib import Path
from typing import Literal
from urllib.parse import urljoin, urlsplit

import anyio
import httpx
from pydantic import BaseModel, Field
from sqlalchemy import func, select

from app.models import (
    Conversation,
    FileLink,
    FileStatus,
    MemoryStatus,
    Message,
    MessagePart,
    StoredFile,
)
from app.services.files import quota
from app.services.files.paths import UnsafePath, safe_filename
from app.services.files.service import storage_for
from app.services.files.sniff import FileRejected, detect
from app.services.jobs import enqueue_job
from app.services.memory.embeddings import Embedder
from app.services.memory.retrieval import prompt_block as memory_prompt_block
from app.services.memory.retrieval import retrieve, usage_part
from app.services.memory.store import add_memory, memory_active_for
from app.services.sync import record_sync
from app.services.tools import workspace as ws
from app.services.tools.netguard import (
    BlockedDestination,
    check_url,
    guarded_client,
    resolve_checked,
)
from app.services.tools.registry import ToolContext, ToolOutput, ToolSpec, builtin
from app.services.tools.sandbox_client import SandboxUnavailable

READ_LIMIT = 200_000
MAX_REDIRECTS = 5
USER_AGENT = "Celius/0.5 (self-hosted assistant; +https://github.com/)"
TEXT_TYPES = (
    "text/",
    "application/json",
    "application/xml",
    "application/xhtml+xml",
    "application/javascript",
    "application/ld+json",
    "application/rss+xml",
    "application/atom+xml",
)


def _error(message: str, **data) -> ToolOutput:
    return ToolOutput(content=f"Error: {message}", data={"error": message, **data}, is_error=True)


def _cap(text: str, limit: int) -> tuple[str, bool]:
    if len(text) <= limit:
        return text, False
    return text[:limit] + f"\n[…cut: {len(text) - limit} more characters]", True


# ---------------------------------------------------------------------- files


class ListFilesArgs(BaseModel):
    path: str = Field(default="", description="Folder inside the workspace; empty for the root")


async def _list_files(ctx: ToolContext, args: ListFilesArgs) -> ToolOutput:
    try:
        entries, more = await anyio.to_thread.run_sync(ws.list_entries, ctx.workspace, args.path)
    except ws.WorkspaceError as exc:
        return _error(str(exc))
    lines = [f"{e.path}/" if e.is_dir else f"{e.path}  ({e.size} bytes)" for e in entries]
    if more:
        lines.append("[…more entries not shown]")
    return ToolOutput(
        content="\n".join(lines) or "(the workspace is empty)",
        data={"entries": [e.__dict__ for e in entries], "truncated": more},
    )


builtin(
    ToolSpec(
        name="list_files",
        label="List files",
        description="List files in this conversation's workspace. Files attached to the "
        "conversation are under attachments/.",
        risk="read",
        handler=_list_files,
        args_model=ListFilesArgs,
    )
)


class ReadFileArgs(BaseModel):
    path: str = Field(description="File path inside the workspace")
    max_chars: int = Field(default=50_000, ge=100, le=READ_LIMIT)


async def _read_file(ctx: ToolContext, args: ReadFileArgs) -> ToolOutput:
    try:
        data, size = await anyio.to_thread.run_sync(
            ws.read_bytes, ctx.workspace, args.path, args.max_chars * 4
        )
    except ws.WorkspaceError as exc:
        return _error(str(exc))
    if b"\x00" in data[:8192]:
        return _error(f"{args.path} is a binary file ({size} bytes); use run_python to inspect it")
    text, cut = _cap(data.decode("utf-8", errors="replace"), args.max_chars)
    return ToolOutput(content=text, data={"path": args.path, "size": size, "truncated": cut})


builtin(
    ToolSpec(
        name="read_file",
        label="Read file",
        description="Read a text file from the workspace.",
        risk="read",
        handler=_read_file,
        args_model=ReadFileArgs,
    )
)


class WriteFileArgs(BaseModel):
    path: str = Field(description="File path inside the workspace")
    content: str = Field(max_length=2_000_000)
    overwrite: bool = True


async def _write_file(ctx: ToolContext, args: WriteFileArgs) -> ToolOutput:
    data = args.content.encode()
    try:
        await anyio.to_thread.run_sync(
            lambda: ws.write_bytes(ctx.workspace, args.path, data, overwrite=args.overwrite)
        )
    except ws.WorkspaceError as exc:
        return _error(str(exc))
    return ToolOutput(
        content=f"Wrote {len(data)} bytes to {args.path}",
        data={"changed_files": [{"path": args.path, "size": len(data)}]},
    )


builtin(
    ToolSpec(
        name="write_file",
        label="Write file",
        description="Create or replace a text file in the workspace.",
        risk="write",
        handler=_write_file,
        args_model=WriteFileArgs,
    )
)


class AttachFileArgs(BaseModel):
    path: str = Field(description="File path inside the workspace")
    name: str | None = Field(default=None, description="File name to show; defaults to the path's")


async def _attach_file(ctx: ToolContext, args: AttachFileArgs) -> ToolOutput:
    settings = ctx.settings
    try:
        filename = safe_filename(args.name or args.path)
    except UnsafePath as exc:
        return _error(str(exc))
    tmp_dir = Path(settings.storage_root) / "tmp"
    tmp_dir.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(dir=tmp_dir)
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        try:
            size = await anyio.to_thread.run_sync(
                ws.copy_out, ctx.workspace, args.path, tmp, settings.max_upload_bytes
            )
            detected = await anyio.to_thread.run_sync(detect, tmp, filename)
        except (ws.WorkspaceError, FileRejected) as exc:
            return _error(str(exc))
        sha256 = await anyio.to_thread.run_sync(_sha256, tmp)
        async with ctx.factory() as db:
            try:
                await quota.reserve(db, ctx.owner_id, size, settings.max_user_storage_bytes)
            except Exception:
                return _error("Not enough storage space left to attach this file")
            await anyio.to_thread.run_sync(storage_for(settings).write_new, tmp, sha256)
            file = StoredFile(
                id=uuid.uuid4(),
                owner_id=ctx.owner_id,
                filename=filename,
                media_type=detected.media_type,
                kind=detected.kind,
                size_bytes=size,
                sha256=sha256,
                status=FileStatus.processing,
                meta={"source": "tool", "workspace_path": args.path},
            )
            db.add(file)
            await db.flush()
            await _link_to_reply(db, ctx.message_id, file)
            enqueue_job(db, ctx.owner_id, "file.process", {"file_id": str(file.id)})
            record_sync(db, ctx.owner_id, "file", file.id)
            await db.commit()
    finally:
        await anyio.Path(tmp).unlink(missing_ok=True)
    return ToolOutput(
        content=f"Attached {filename} ({size} bytes) to the reply.",
        data={"file_id": str(file.id), "filename": filename, "size": size},
    )


async def _link_to_reply(db, message_id: uuid.UUID, file: StoredFile) -> None:
    """Show the file on the reply, like an attachment on a user message."""
    message = await db.get(Message, message_id)
    link_position = await db.scalar(
        select(func.max(FileLink.position)).where(FileLink.message_id == message_id)
    )
    db.add(FileLink(message_id=message_id, file_id=file.id, position=(link_position or 0) + 1))
    part_position = await db.scalar(
        select(func.max(MessagePart.position)).where(
            MessagePart.message_id == message_id,
            MessagePart.revision == message.current_revision,
        )
    )
    db.add(
        MessagePart(
            message_id=message_id,
            revision=message.current_revision,
            position=(part_position if part_position is not None else -1) + 1,
            type="file",
            data={
                "file_id": str(file.id),
                "filename": file.filename,
                "relative_path": None,
                "media_type": file.media_type,
                "kind": file.kind.value,
                "size_bytes": file.size_bytes,
                "has_thumbnail": False,
                "source": "tool",
            },
        )
    )


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


builtin(
    ToolSpec(
        name="attach_file",
        label="Attach file",
        description="Attach a workspace file to your reply so the user can download it. Use this "
        "for every file you produced that the user should get.",
        risk="write",
        handler=_attach_file,
        args_model=AttachFileArgs,
    )
)


# ---------------------------------------------------------------------- sandbox


def _sandbox_output(result: dict, ctx: ToolContext) -> ToolOutput:
    parts = []
    if result.get("stdout"):
        parts.append(
            result["stdout"] + ("\n[…output cut]" if result.get("stdout_truncated") else "")
        )
    if result.get("stderr"):
        cut = "\n[…cut]" if result.get("stderr_truncated") else ""
        parts.append("stderr:\n" + result["stderr"] + cut)
    reason = result.get("killed_reason")
    if reason:
        parts.append(f"[stopped: {reason.replace('_', ' ')}]")
    parts.append(f"[exit code {result.get('exit_code')}]")
    changed = result.get("changed_files") or []
    if changed:
        listed = ", ".join(f"{c['path']} ({c['size']} bytes)" for c in changed[:20])
        parts.append(f"[files changed: {listed}]")
    if result.get("workspace_over_limit"):
        parts.append("[warning: the workspace is over its size limit; delete files]")
    failed = reason is not None or (result.get("exit_code") or 0) != 0
    return ToolOutput(content="\n".join(parts), data=result, is_error=failed)


async def _execute(ctx: ToolContext, kind: str, code: str, seconds: int | None) -> ToolOutput:
    if ctx.sandbox is None:
        return _error("The sandbox is not configured")
    limit = min(seconds or ctx.settings.tool_call_timeout_seconds, ctx.time_limit)
    try:
        result = await ctx.sandbox.execute(
            str(ctx.conversation_id),
            kind,
            code,
            timeout_seconds=limit,
            memory_mb=ctx.settings.tool_memory_mb,
            output_limit=min(ctx.settings.tool_max_result_chars * 2, 1024 * 1024),
        )
    except SandboxUnavailable as exc:
        return _error(str(exc))
    return _sandbox_output(result, ctx)


class PythonArgs(BaseModel):
    code: str = Field(max_length=200_000, description="Python 3 source to run")
    timeout_seconds: int | None = Field(default=None, ge=1, le=3600)


async def _run_python(ctx: ToolContext, args: PythonArgs) -> ToolOutput:
    return await _execute(ctx, "python", args.code, args.timeout_seconds)


builtin(
    ToolSpec(
        name="run_python",
        label="Run Python",
        description="Run Python 3 in an isolated sandbox with no network. The working directory "
        "is the conversation's workspace and files persist between calls. numpy, pandas, "
        "matplotlib (save figures to files), Pillow, python-docx, openpyxl, python-pptx and "
        "fpdf2 are installed; nothing else can be installed. Print what you need to see.",
        risk="execute",
        handler=_run_python,
        args_model=PythonArgs,
        needs_sandbox=True,
    )
)


class ShellArgs(BaseModel):
    command: str = Field(max_length=100_000, description="bash commands to run")
    timeout_seconds: int | None = Field(default=None, ge=1, le=3600)


async def _run_shell(ctx: ToolContext, args: ShellArgs) -> ToolOutput:
    return await _execute(ctx, "shell", args.command, args.timeout_seconds)


builtin(
    ToolSpec(
        name="run_shell",
        label="Run shell",
        description="Run bash commands in the same sandbox and workspace as run_python "
        "(no network, no root, standard Unix tools and git).",
        risk="execute",
        handler=_run_shell,
        args_model=ShellArgs,
        needs_sandbox=True,
    )
)


GIT_NETWORK_COMMANDS = {"clone", "fetch", "pull", "push", "ls-remote", "submodule", "remote"}


class GitArgs(BaseModel):
    args: list[str] = Field(
        min_length=1, max_length=50, description='git arguments, e.g. ["log", "--oneline", "-5"]'
    )
    repository: str = Field(default="", description="Repository folder in the workspace")


async def _git(ctx: ToolContext, args: GitArgs) -> ToolOutput:
    subcommand = next((a for a in args.args if not a.startswith("-")), "")
    if subcommand in GIT_NETWORK_COMMANDS:
        return _error(f"git {subcommand} needs the network; use git_clone to get a repository")
    try:
        ws.resolve(ctx.workspace, args.repository, allow_root=True)
    except ws.WorkspaceError as exc:
        return _error(str(exc))
    command = shlex.join(["git", "-C", args.repository or ".", *args.args])
    preamble = (
        "export GIT_TERMINAL_PROMPT=0 GIT_CONFIG_NOSYSTEM=1 GIT_AUTHOR_NAME=Celius "
        "GIT_AUTHOR_EMAIL=celius@localhost GIT_COMMITTER_NAME=Celius "
        "GIT_COMMITTER_EMAIL=celius@localhost\n"
    )
    return await _execute(ctx, "shell", preamble + command, None)


builtin(
    ToolSpec(
        name="git",
        label="Git",
        description="Run a git command on a repository in the workspace (status, log, diff, "
        "add, commit, branch, checkout, init…). No network access.",
        risk="execute",
        handler=_git,
        args_model=GitArgs,
        needs_sandbox=True,
    )
)


class GitCloneArgs(BaseModel):
    url: str = Field(description="https URL of a public Git repository")
    directory: str | None = Field(default=None, description="Target folder in the workspace")


async def _git_clone(ctx: ToolContext, args: GitCloneArgs) -> ToolOutput:
    """The worker downloads (no checkout, so nothing from the repository is written as
    files here); the sandbox then checks the files out."""
    if not args.url.startswith("https://"):
        return _error("Only https:// repository URLs are allowed")
    try:
        host = check_url(args.url, ctx.settings.web_allowlist)
        addresses = await resolve_checked(host, 443, resolver=ctx.resolver)
    except BlockedDestination as exc:
        return _error(str(exc))
    name = args.directory or urlsplit(args.url).path.rstrip("/").rsplit("/", 1)[-1]
    name = name.removesuffix(".git") or "repository"
    try:
        target = ws.resolve(ctx.workspace, name)
    except ws.WorkspaceError as exc:
        return _error(str(exc))
    if await anyio.to_thread.run_sync(os.path.lexists, target):
        return _error(f"{name} already exists in the workspace")
    pin = ",".join(f"[{a}]" if ":" in a else a for a in addresses)
    env = {
        "PATH": os.environ.get("PATH", "/usr/bin:/bin"),
        "HOME": tempfile.gettempdir(),
        "GIT_TERMINAL_PROMPT": "0",
        "GIT_CONFIG_NOSYSTEM": "1",
        "GIT_CONFIG_GLOBAL": "/dev/null",
        "GIT_ALLOW_PROTOCOL": "https",
    }
    command = [
        "git",
        "-c",
        f"http.curloptResolve={host}:443:{pin}",  # connect only to the checked address
        "-c",
        "http.followRedirects=false",
        "-c",
        "protocol.allow=never",
        "-c",
        "protocol.https.allow=always",
        "-c",
        "core.hooksPath=/dev/null",
        "clone",
        "--no-checkout",
        "--depth",
        "50",
        "--no-recurse-submodules",
        "--quiet",
        "--",
        args.url,
        str(target),
    ]
    try:
        process = await asyncio.create_subprocess_exec(
            *command, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE, env=env
        )
    except FileNotFoundError:
        return _error("git is not installed on the worker")
    try:
        _, stderr = await asyncio.wait_for(process.communicate(), timeout=ctx.time_limit)
    except (TimeoutError, asyncio.CancelledError):
        process.kill()
        await process.wait()
        await anyio.to_thread.run_sync(lambda: _remove_tree(target))
        raise
    if process.returncode != 0:
        await anyio.to_thread.run_sync(lambda: _remove_tree(target))
        message = stderr.decode(errors="replace").strip()[-500:]
        return _error(f"git clone failed: {message}")
    checkout = await _execute(
        ctx, "shell", shlex.join(["git", "-C", name, "reset", "-q", "--hard", "HEAD"]), None
    )
    if checkout.is_error:
        return checkout
    return ToolOutput(content=f"Cloned {args.url} into {name}/", data={"path": name})


def _remove_tree(path: Path) -> None:
    shutil.rmtree(path, ignore_errors=True)


builtin(
    ToolSpec(
        name="git_clone",
        label="Clone repository",
        description="Clone a public Git repository over https into the workspace (latest 50 "
        "commits). Then use git, read_file or run_shell on it.",
        risk="network",
        handler=_git_clone,
        args_model=GitCloneArgs,
        needs_sandbox=True,
        needs_network=True,
    )
)


# ---------------------------------------------------------------------- web


class _TextExtractor(HTMLParser):
    SKIP = {"script", "style", "noscript", "template", "svg", "head"}
    BLOCK = {
        "p",
        "div",
        "br",
        "li",
        "tr",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "section",
        "article",
        "header",
        "footer",
        "pre",
        "blockquote",
        "table",
    }

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.skipping = 0
        self.title = ""
        self._in_title = False

    def handle_starttag(self, tag, attrs):
        if tag == "title":
            self._in_title = True
        if tag in self.SKIP:
            self.skipping += 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag == "title":
            self._in_title = False
        if tag in self.SKIP and self.skipping:
            self.skipping -= 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        elif not self.skipping:
            self.parts.append(data)

    def text(self) -> str:
        lines = (" ".join(line.split()) for line in "".join(self.parts).splitlines())
        out, blank = [], False
        for line in lines:
            if line:
                out.append(line)
                blank = False
            elif not blank and out:
                out.append("")
                blank = True
        return "\n".join(out).strip()


def html_to_text(markup: str) -> tuple[str, str]:
    parser = _TextExtractor()
    parser.feed(markup)
    parser.close()
    return html.unescape(parser.title.strip()), parser.text()


class WebFetchArgs(BaseModel):
    url: str = Field(description="http(s) URL to fetch")
    save_as: str | None = Field(
        default=None,
        description="Save the response body to this workspace path instead of returning "
        "its text (use for data files and images; the sandbox itself has no network).",
    )


async def fetch(ctx: ToolContext, url: str, *, transport=None) -> tuple[httpx.Response, bytes, str]:
    """GET with the SSRF guard, re-checking every redirect. Returns the final response,
    its (capped) body and the final URL."""
    limit = ctx.settings.tool_web_max_bytes
    async with guarded_client(
        timeout=min(ctx.time_limit, 60),
        resolver=ctx.resolver,
        transport=transport,
        headers={
            "User-Agent": USER_AGENT,
            "Accept": "text/html,application/json,text/*;q=0.9,*/*;q=0.5",
        },
    ) as client:
        for _ in range(MAX_REDIRECTS + 1):
            check_url(url, ctx.settings.web_allowlist)
            async with client.stream("GET", url) as response:
                if response.is_redirect and "location" in response.headers:
                    url = urljoin(url, response.headers["location"])
                    continue
                body = bytearray()
                async for chunk in response.aiter_bytes():
                    body.extend(chunk)
                    if len(body) > limit:
                        raise BlockedDestination(f"The response is larger than {limit} bytes")
                return response, bytes(body), url
    raise BlockedDestination("Too many redirects")


async def _web_fetch(ctx: ToolContext, args: WebFetchArgs) -> ToolOutput:
    try:
        response, body, final_url = await fetch(ctx, args.url, transport=ctx.web_transport)
    except BlockedDestination as exc:
        return _error(str(exc), url=args.url)
    except httpx.HTTPError as exc:
        return _error(f"The request failed: {type(exc).__name__}", url=args.url)
    content_type = response.headers.get("content-type", "").split(";")[0].strip().lower()
    data = {
        "url": final_url,
        "status": response.status_code,
        "content_type": content_type,
        "bytes": len(body),
    }
    if args.save_as:
        try:
            await anyio.to_thread.run_sync(
                lambda: ws.write_bytes(ctx.workspace, args.save_as, body, overwrite=True)
            )
        except ws.WorkspaceError as exc:
            return _error(str(exc), **data)
        return ToolOutput(
            content=f"HTTP {response.status_code}; saved {len(body)} bytes ({content_type}) "
            f"to {args.save_as}",
            data={**data, "changed_files": [{"path": args.save_as, "size": len(body)}]},
            is_error=response.status_code >= 400,
        )
    if not content_type.startswith(TEXT_TYPES) and content_type:
        return _error(f"The response is {content_type}, not text; pass save_as to store it", **data)
    text = body.decode(response.encoding or "utf-8", errors="replace")
    title = ""
    if content_type in ("text/html", "application/xhtml+xml") or text.lstrip()[:15].lower() in (
        "<!doctype html>",
    ):
        title, text = html_to_text(text)
    text, cut = _cap(text, ctx.settings.tool_max_result_chars)
    header = f"URL: {final_url}\nHTTP {response.status_code}" + (
        f"\nTitle: {title}" if title else ""
    )
    return ToolOutput(
        content=f"{header}\n\n{text}",
        data={**data, "title": title, "truncated": cut},
        is_error=response.status_code >= 400,
    )


builtin(
    ToolSpec(
        name="web_fetch",
        label="Fetch web page",
        description="Fetch a public web page or file over http(s). Returns readable text for "
        "HTML and text; use save_as to store other content in the workspace.",
        risk="network",
        handler=_web_fetch,
        args_model=WebFetchArgs,
        needs_network=True,
    )
)


# ---------------------------------------------------------------------- images and documents


def _script(body: str, payload: dict) -> str:
    """A fixed script with its arguments as JSON data (never interpolated as code)."""
    return f"import json\nARGS = json.loads({json.dumps(json.dumps(payload))})\n{body}"


IMAGE_SCRIPT = r"""
from PIL import Image, ImageOps
src, op, out = ARGS["path"], ARGS["operation"], ARGS.get("output_path")
img = Image.open(src)
img = ImageOps.exif_transpose(img)
if op == "info":
    print(json.dumps({"format": img.format, "mode": img.mode, "width": img.width,
                      "height": img.height}))
    raise SystemExit(0)
if op == "resize":
    w, h = ARGS.get("width"), ARGS.get("height")
    if w and h:
        img = img.resize((w, h))
    else:
        img.thumbnail((w or 100000, h or 100000))
elif op == "crop":
    img = img.crop(tuple(ARGS["box"]))
elif op == "rotate":
    img = img.rotate(-ARGS.get("degrees", 90), expand=True)
elif op == "grayscale":
    img = ImageOps.grayscale(img)
elif op == "convert":
    pass
fmt = (out.rsplit(".", 1)[-1] if "." in out else "png").upper().replace("JPG", "JPEG")
if fmt == "JPEG" and img.mode not in ("RGB", "L"):
    img = img.convert("RGB")
img.save(out, fmt)
print(f"saved {out} ({img.width}x{img.height})")
"""


class ImageArgs(BaseModel):
    path: str = Field(description="Image file in the workspace")
    operation: Literal["info", "resize", "crop", "rotate", "grayscale", "convert"]
    output_path: str | None = Field(default=None, description="Where to save the result")
    width: int | None = Field(default=None, ge=1, le=20000)
    height: int | None = Field(default=None, ge=1, le=20000)
    box: list[int] | None = Field(
        default=None, min_length=4, max_length=4, description="Crop box: left, top, right, bottom"
    )
    degrees: int | None = Field(default=None, ge=-360, le=360)


async def _image(ctx: ToolContext, args: ImageArgs) -> ToolOutput:
    try:
        ws.resolve(ctx.workspace, args.path)
        if args.operation != "info":
            if not args.output_path:
                return _error("output_path is required for this operation")
            ws.resolve(ctx.workspace, args.output_path)
    except ws.WorkspaceError as exc:
        return _error(str(exc))
    if args.operation == "crop" and not args.box:
        return _error("box is required to crop")
    return await _execute(ctx, "python", _script(IMAGE_SCRIPT, args.model_dump()), None)


builtin(
    ToolSpec(
        name="image",
        label="Edit image",
        description="Inspect or transform an image in the workspace: info, resize, crop, "
        "rotate, grayscale, or convert (format from output_path's extension).",
        risk="write",
        handler=_image,
        args_model=ImageArgs,
        needs_sandbox=True,
    )
)


DOCUMENT_SCRIPT = r"""
import csv, re
fmt, out = ARGS["format"], ARGS["path"]
title, md = ARGS.get("title"), ARGS.get("markdown") or ""
sheets, slides = ARGS.get("sheets") or [], ARGS.get("slides") or []

def blocks(text):
    for line in text.splitlines():
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            yield "h", len(m.group(1)), m.group(2)
        elif re.match(r"^\s*[-*]\s+", line):
            yield "li", 0, re.sub(r"^\s*[-*]\s+", "", line)
        elif line.strip():
            yield "p", 0, line.strip()

if fmt == "docx":
    from docx import Document
    doc = Document()
    if title:
        doc.add_heading(title, 0)
    for kind, level, text in blocks(md):
        if kind == "h":
            doc.add_heading(text, min(level, 9))
        elif kind == "li":
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)
    doc.save(out)
elif fmt == "pdf":
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.add_font("DejaVu", "", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    pdf.add_font("DejaVu", "B", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")
    if title:
        pdf.set_font("DejaVu", "B", 18)
        pdf.multi_cell(0, 10, title)
        pdf.ln(2)
    for kind, level, text in blocks(md):
        if kind == "h":
            pdf.set_font("DejaVu", "B", max(18 - 2 * level, 11))
            pdf.multi_cell(0, 8, text)
        else:
            pdf.set_font("DejaVu", "", 11)
            pdf.multi_cell(0, 6, ("• " if kind == "li" else "") + text)
    pdf.output(out)
elif fmt == "xlsx":
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in sheets or [{"name": "Sheet1", "rows": []}]:
        ws = wb.create_sheet(str(sheet.get("name") or "Sheet")[:31])
        for row in sheet.get("rows") or []:
            ws.append(row)
    wb.save(out)
elif fmt == "csv":
    rows = (sheets[0].get("rows") if sheets else []) or []
    with open(out, "w", newline="") as handle:
        csv.writer(handle).writerows(rows)
elif fmt == "pptx":
    from pptx import Presentation
    deck = Presentation()
    if title:
        slide = deck.slides.add_slide(deck.slide_layouts[0])
        slide.shapes.title.text = title
    for item in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = item.get("title") or ""
        body = slide.placeholders[1].text_frame
        for i, bullet in enumerate(item.get("bullets") or []):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = str(bullet)
    deck.save(out)
elif fmt in ("md", "html"):
    text = (f"# {title}\n\n" if title else "") + md
    if fmt == "html":
        import html as h
        parts = []
        for kind, level, line in blocks(text):
            tag = f"h{level}" if kind == "h" else kind
            parts.append(f"<{tag}>{h.escape(line)}</{tag}>")
        text = "<!doctype html><meta charset=utf-8>\n" + "\n".join(parts)
    with open(out, "w") as handle:
        handle.write(text)
import os
print(f"saved {out} ({os.path.getsize(out)} bytes)")
"""


class SheetArg(BaseModel):
    name: str = "Sheet1"
    rows: list[list[str | int | float | bool | None]] = Field(default_factory=list)


class SlideArg(BaseModel):
    title: str = ""
    bullets: list[str] = Field(default_factory=list)


class DocumentArgs(BaseModel):
    format: Literal["docx", "pdf", "xlsx", "csv", "pptx", "md", "html"]
    path: str = Field(description="Output file path in the workspace")
    title: str | None = None
    markdown: str | None = Field(
        default=None,
        max_length=500_000,
        description="Body for docx/pdf/md/html: # headings, - bullets, paragraphs",
    )
    sheets: list[SheetArg] | None = Field(default=None, description="For xlsx/csv")
    slides: list[SlideArg] | None = Field(default=None, description="For pptx")


async def _make_document(ctx: ToolContext, args: DocumentArgs) -> ToolOutput:
    try:
        ws.resolve(ctx.workspace, args.path)
    except ws.WorkspaceError as exc:
        return _error(str(exc))
    return await _execute(ctx, "python", _script(DOCUMENT_SCRIPT, args.model_dump()), None)


builtin(
    ToolSpec(
        name="make_document",
        label="Make document",
        description="Create a Word, PDF, Excel, CSV, PowerPoint, Markdown or HTML file in the "
        "workspace from structured content. Then attach_file it for the user.",
        risk="write",
        handler=_make_document,
        args_model=DocumentArgs,
        needs_sandbox=True,
    )
)


# ---------------------------------------------------------------------- memory


class RememberArgs(BaseModel):
    fact: str = Field(
        min_length=3,
        max_length=1000,
        description="One durable fact about the user, in the third person",
    )
    scope: Literal["global", "project"] = Field(
        default="global", description="project: only for chats in this chat's project"
    )


async def _remember(ctx: ToolContext, args: RememberArgs) -> ToolOutput:
    async with ctx.factory() as db:
        conversation = await db.get(Conversation, ctx.conversation_id)
        if conversation is None or not await memory_active_for(db, conversation, ctx.settings):
            return _error("Memory is turned off for this chat")
        project_id = conversation.project_id if args.scope == "project" else None
        embedder = Embedder.from_settings(ctx.settings)
        try:
            memory = await add_memory(
                db,
                ctx.owner_id,
                args.fact,
                ctx.settings,
                source="model",
                scope="project" if project_id else "global",
                project_id=project_id,
                conversation_id=ctx.conversation_id,
                message_id=ctx.message_id,
                embedder=embedder,
            )
        finally:
            if embedder is not None:
                await embedder.aclose()
        await db.commit()
    if memory is None:
        return ToolOutput(content="Already remembered.", data={"duplicate": True})
    if memory.status == MemoryStatus.proposed:
        return ToolOutput(
            content="Proposed; the user will review it before it is remembered.",
            data={"memory_id": str(memory.id), "status": "proposed"},
        )
    return ToolOutput(content="Remembered.", data={"memory_id": str(memory.id), "status": "active"})


builtin(
    ToolSpec(
        name="remember",
        label="Remember",
        description="Propose a lasting fact about the user (a preference, their setup, an "
        "ongoing project) to remember in future chats. Only when they would clearly want it "
        "kept; never secrets or sensitive details.",
        risk="read",
        handler=_remember,
        args_model=RememberArgs,
    )
)


class SearchMemoryArgs(BaseModel):
    query: str = Field(min_length=2, max_length=500, description="What to look for")


async def _search_memory(ctx: ToolContext, args: SearchMemoryArgs) -> ToolOutput:
    async with ctx.factory() as db:
        conversation = await db.get(Conversation, ctx.conversation_id)
        if conversation is None:
            return _error("Conversation not found")
        embedder = Embedder.from_settings(ctx.settings)
        try:
            found = await retrieve(db, conversation, args.query, ctx.settings, embedder)
        finally:
            if embedder is not None:
                await embedder.aclose()
        await db.commit()
    if found.empty:
        return ToolOutput(content="Nothing found in memory.", data={"results": 0})
    block = memory_prompt_block(found, ctx.settings.tool_max_result_chars)
    return ToolOutput(content=block, data=usage_part(found))


builtin(
    ToolSpec(
        name="search_memory",
        label="Search memory",
        description="Search what you remember about the user and their earlier chats (in the "
        "same project). Use when they refer to something from another conversation.",
        risk="read",
        handler=_search_memory,
        args_model=SearchMemoryArgs,
    )
)

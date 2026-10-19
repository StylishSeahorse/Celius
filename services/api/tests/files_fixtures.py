"""Real files of each supported type, built in memory for tests."""

import io
import stat
import zipfile


def png(width=64, height=48, color=(200, 80, 40)) -> bytes:
    from PIL import Image

    out = io.BytesIO()
    Image.new("RGB", (width, height), color).save(out, "PNG")
    return out.getvalue()


def pdf(text="Hello PDF world") -> bytes:
    """A minimal one-page PDF with real text (pypdf can extract it)."""
    stream = f"BT /F1 18 Tf 20 100 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def docx(text="Hello from Word") -> bytes:
    import docx as python_docx

    document = python_docx.Document()
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "cell A"
    table.rows[0].cells[1].text = "cell B"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def xlsx() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet.append(["item", "cost"])
    sheet.append(["seeds", 12])
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def pptx(text="Hello from slides") -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = text
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue()


def zip_of(entries: dict[str, bytes], *, symlinks: dict[str, str] | None = None) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
        for name, target in (symlinks or {}).items():
            info = zipfile.ZipInfo(name)
            info.external_attr = (stat.S_IFLNK | 0o777) << 16
            archive.writestr(info, target)
    return out.getvalue()
